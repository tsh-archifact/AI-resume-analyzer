"""
DOCX Resume Generator Service.
Provides 5 professional, ATS-compliant templates with custom typography,
accent palettes, layout spacing, and markdown-to-DOCX conversion.
"""

from __future__ import annotations

import re
import tempfile
from dataclasses import dataclass, field
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


@dataclass(frozen=True)
class DocxTemplateStyle:
    template_id: str
    name: str
    description: str
    persona: str
    body_font: str
    heading_font: str
    accent_hex: str
    secondary_hex: str
    body_size_pt: float = 10.5
    title_size_pt: float = 20.0
    h2_size_pt: float = 12.0
    h3_size_pt: float = 11.0
    line_spacing: float = 1.15
    space_after_body_pt: float = 4.0
    space_before_h2_pt: float = 10.0
    space_after_h2_pt: float = 3.0
    center_header: bool = False
    all_caps_h2: bool = True
    add_heading_border: bool = True
    border_size: str = "12"  # 1/8 pt units (12 = 1.5pt)
    margin_inches: float = 0.75
    features: list[str] = field(default_factory=list)

    @property
    def accent_rgb(self) -> RGBColor:
        h = self.accent_hex.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    @property
    def secondary_rgb(self) -> RGBColor:
        h = self.secondary_hex.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ---------------------------------------------------------------------------
# 5 Professional Resume Templates
# ---------------------------------------------------------------------------

TEMPLATES: dict[str, DocxTemplateStyle] = {
    "modern_teal": DocxTemplateStyle(
        template_id="modern_teal",
        name="Modern Minimalist",
        description="Clean, contemporary sans-serif layout with deep teal accents. Ideal for tech, product, and modern businesses.",
        persona="Recommended for Tech, Product, & Startups",
        body_font="Calibri",
        heading_font="Calibri",
        accent_hex="0F766E",  # Deep Teal
        secondary_hex="475569",  # Slate
        body_size_pt=10.5,
        title_size_pt=20.0,
        h2_size_pt=12.5,
        h3_size_pt=11.0,
        line_spacing=1.15,
        space_after_body_pt=4.0,
        space_before_h2_pt=9.0,
        space_after_h2_pt=3.0,
        center_header=False,
        all_caps_h2=True,
        add_heading_border=True,
        border_size="10",
        margin_inches=0.75,
        features=["Teal section underline", "Calibri clean sans", "Left-aligned modern header"],
    ),
    "executive_navy": DocxTemplateStyle(
        template_id="executive_navy",
        name="Executive Classic",
        description="Authoritative serif typography in rich navy blue. Engineered for senior executives, finance, consulting, and legal roles.",
        persona="Recommended for Leadership, Finance, & Consulting",
        body_font="Georgia",
        heading_font="Georgia",
        accent_hex="1E3A8A",  # Deep Navy Blue
        secondary_hex="334155",  # Dark Slate
        body_size_pt=10.0,
        title_size_pt=22.0,
        h2_size_pt=12.0,
        h3_size_pt=10.5,
        line_spacing=1.2,
        space_after_body_pt=4.5,
        space_before_h2_pt=11.0,
        space_after_h2_pt=4.0,
        center_header=True,
        all_caps_h2=True,
        add_heading_border=True,
        border_size="14",
        margin_inches=0.8,
        features=["Centered prestige header", "Georgia elegant serif", "Navy section dividers"],
    ),
    "tech_indigo": DocxTemplateStyle(
        template_id="tech_indigo",
        name="Tech & Developer",
        description="Modern high-clarity typography with electric indigo accents and structured sections. Tailored for software engineers and DevOps.",
        persona="Recommended for Software Engineers & DevOps",
        body_font="Segoe UI",
        heading_font="Segoe UI Semibold",
        accent_hex="4338CA",  # Indigo
        secondary_hex="374151",  # Neutral dark
        body_size_pt=10.0,
        title_size_pt=20.0,
        h2_size_pt=12.0,
        h3_size_pt=10.5,
        line_spacing=1.15,
        space_after_body_pt=3.5,
        space_before_h2_pt=9.0,
        space_after_h2_pt=3.0,
        center_header=False,
        all_caps_h2=True,
        add_heading_border=True,
        border_size="12",
        margin_inches=0.7,
        features=["Indigo accent line", "Segoe UI modern tech font", "Clean compact bullets"],
    ),
    "elegant_burgundy": DocxTemplateStyle(
        template_id="elegant_burgundy",
        name="Elegant Academic",
        description="Refined and distinguished serif aesthetic with deep burgundy accents. Perfect for academia, research, medicine, and writing.",
        persona="Recommended for Academia, Research, & Medical",
        body_font="Cambria",
        heading_font="Cambria",
        accent_hex="881337",  # Deep Burgundy
        secondary_hex="4A5568",  # Charcoal
        body_size_pt=10.5,
        title_size_pt=21.0,
        h2_size_pt=12.5,
        h3_size_pt=11.0,
        line_spacing=1.2,
        space_after_body_pt=4.0,
        space_before_h2_pt=10.0,
        space_after_h2_pt=3.5,
        center_header=True,
        all_caps_h2=True,
        add_heading_border=True,
        border_size="10",
        margin_inches=0.8,
        features=["Centered title", "Deep burgundy styling", "Cambria academic serif"],
    ),
    "compact_slate": DocxTemplateStyle(
        template_id="compact_slate",
        name="Compact High-Density",
        description="Space-efficient, high-density layout designed to fit extensive career histories into fewer pages without sacrificing readability.",
        persona="Recommended for Multi-Page Content / Maximum Space",
        body_font="Arial",
        heading_font="Arial",
        accent_hex="1F2937",  # Deep Graphite
        secondary_hex="4B5563",  # Medium Gray
        body_size_pt=9.5,
        title_size_pt=18.0,
        h2_size_pt=11.0,
        h3_size_pt=10.0,
        line_spacing=1.1,
        space_after_body_pt=2.5,
        space_before_h2_pt=7.0,
        space_after_h2_pt=2.0,
        center_header=False,
        all_caps_h2=True,
        add_heading_border=True,
        border_size="8",
        margin_inches=0.55,
        features=["Slim 0.55-in margins", "High data density", "Graphite ATS-optimized"],
    ),
}

DEFAULT_TEMPLATE_ID = "modern_teal"


def get_available_templates() -> list[dict[str, Any]]:
    """Return public metadata for all available DOCX resume templates."""
    return [
        {
            "template_id": style.template_id,
            "name": style.name,
            "description": style.description,
            "persona": style.persona,
            "font_family": style.body_font,
            "accent_hex": style.accent_hex,
            "secondary_hex": style.secondary_hex,
            "features": style.features,
            "is_default": style.template_id == DEFAULT_TEMPLATE_ID,
        }
        for style in TEMPLATES.values()
    ]


# ---------------------------------------------------------------------------
# XML / Word Styling Helpers
# ---------------------------------------------------------------------------

def _add_bottom_border(paragraph, color_hex: str, size: str = "12") -> None:
    """Add a sleek bottom accent border to a paragraph in Word XML."""
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    pBdr.append(bottom)
    pPr.append(pBdr)


def _append_markdown_inline_runs(paragraph, text: str, font_name: str, font_size_pt: float, default_color: RGBColor | None = None) -> None:
    """Parse inline markdown **bold** and *italic* tokens and add styled runs."""
    # Split text by bold markers (**bold**)
    tokens = re.split(r"(\*\*.*?\*\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            clean_text = token[2:-2]
            run = paragraph.add_run(clean_text)
            run.bold = True
        else:
            # Check for italics inside non-bold text
            sub_tokens = re.split(r"(\*.*?\*)", token)
            for sub_token in sub_tokens:
                if not sub_token:
                    continue
                if sub_token.startswith("*") and sub_token.endswith("*") and len(sub_token) >= 3:
                    run = paragraph.add_run(sub_token[1:-1])
                    run.italic = True
                else:
                    run = paragraph.add_run(sub_token)

        if run:
            run.font.name = font_name
            run.font.size = Pt(font_size_pt)
            if default_color:
                run.font.color.rgb = default_color


# Known standard resume section headers
_KNOWN_HEADINGS = {
    "EXPERIENCE", "WORK EXPERIENCE", "PROFESSIONAL EXPERIENCE", "EMPLOYMENT HISTORY",
    "EDUCATION", "ACADEMIC BACKGROUND", "ACADEMIC HISTORY",
    "SKILLS", "TECHNICAL SKILLS", "CORE COMPETENCIES", "KEY SKILLS", "AREAS OF EXPERTISE",
    "PROJECTS", "KEY PROJECTS", "PERSONAL PROJECTS", "ACADEMIC PROJECTS",
    "CERTIFICATIONS", "CERTIFICATES", "LICENSES & CERTIFICATIONS",
    "SUMMARY", "PROFESSIONAL SUMMARY", "EXECUTIVE SUMMARY", "CAREER OBJECTIVE", "OBJECTIVE", "PROFILE",
    "AWARDS", "HONORS & AWARDS", "ACHIEVEMENTS",
    "PUBLICATIONS", "VOLUNTEER WORK", "VOLUNTEERING", "LANGUAGES", "INTERESTS", "REFERENCES",
}


def _is_section_heading(line: str) -> tuple[bool, str]:
    """Check if line is a markdown ## heading or ALL-CAPS resume section."""
    stripped = line.strip()
    if not stripped:
        return False, ""

    # Markdown ## or ### heading
    if re.match(r"^#{2,3}\s+", stripped):
        clean = re.sub(r"^#{2,3}\s+", "", stripped).strip()
        return True, clean

    # Check uppercase resume headings
    normalized = stripped.rstrip(":").upper()
    if normalized in _KNOWN_HEADINGS:
        return True, stripped.rstrip(":")

    words = stripped.split()
    if 1 <= len(words) <= 6:
        alpha_chars = [c for c in stripped if c.isalpha()]
        if len(alpha_chars) >= 3 and all(c.isupper() for c in alpha_chars):
            return True, stripped.rstrip(":")

    return False, ""


def _is_bullet_point(line: str) -> tuple[bool, str]:
    """Check if a line represents a bullet point."""
    stripped = line.strip()
    bullet_match = re.match(r"^[-*•–—]\s+(.+)$", stripped)
    if bullet_match:
        return True, bullet_match.group(1).strip()
    return False, ""


def _is_contact_line(line: str) -> bool:
    """Identify contact detail lines (email, phone, location, linkedin)."""
    stripped = line.strip()
    return bool(
        "@" in stripped or
        re.search(r"\b(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b", stripped) or
        "linkedin.com" in stripped.lower() or
        "github.com" in stripped.lower()
    )


# ---------------------------------------------------------------------------
# Generator Function
# ---------------------------------------------------------------------------

def generate_resume_docx(resume_text: str, template_id: str = DEFAULT_TEMPLATE_ID) -> str:
    """Generate a high-quality, ATS-compliant DOCX from resume markdown or text."""
    style = TEMPLATES.get(template_id, TEMPLATES[DEFAULT_TEMPLATE_ID])
    document = Document()

    # --- Set Margins ---
    for section in document.sections:
        section.top_margin = Inches(style.margin_inches)
        section.bottom_margin = Inches(style.margin_inches)
        section.left_margin = Inches(style.margin_inches)
        section.right_margin = Inches(style.margin_inches)

    # --- Setup Base Styles ---
    normal_style = document.styles["Normal"]
    normal_style.font.name = style.body_font
    normal_style.font.size = Pt(style.body_size_pt)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)  # Dark charcoal text

    try:
        h2_style = document.styles["Heading 2"]
        h2_style.font.name = style.heading_font
        h2_style.font.size = Pt(style.h2_size_pt)
        h2_style.font.bold = True
        h2_style.font.color.rgb = style.accent_rgb
    except Exception:
        pass

    raw_lines = [line.strip() for line in resume_text.split("\n")]
    # Filter out consecutive empty lines
    lines: list[str] = []
    for line in raw_lines:
        if line or (lines and lines[-1]):
            lines.append(line)

    is_first_non_empty = True
    idx = 0

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        if not stripped:
            idx += 1
            continue

        # 1. Candidate Name (First Line or Markdown # Title)
        if is_first_non_empty:
            is_first_non_empty = False
            candidate_name = re.sub(r"^#\s*", "", stripped).strip()
            title_para = document.add_paragraph()
            title_para.paragraph_format.space_before = Pt(0)
            title_para.paragraph_format.space_after = Pt(2)
            if style.center_header:
                title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = title_para.add_run(candidate_name)
            run.font.name = style.heading_font
            run.font.size = Pt(style.title_size_pt)
            run.font.bold = True
            run.font.color.rgb = style.accent_rgb
            idx += 1
            continue

        # 2. Contact Information / Subtitle Line (immediately after name)
        if _is_contact_line(stripped) or (idx <= 3 and ("|" in stripped or "•" in stripped)):
            contact_para = document.add_paragraph()
            contact_para.paragraph_format.space_before = Pt(0)
            contact_para.paragraph_format.space_after = Pt(style.space_before_h2_pt)
            if style.center_header:
                contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            _append_markdown_inline_runs(
                contact_para,
                stripped,
                style.body_font,
                style.body_size_pt - 0.5,
                style.secondary_rgb,
            )
            idx += 1
            continue

        # 3. Section Headings (## Section or ALL-CAPS)
        is_heading, heading_title = _is_section_heading(stripped)
        if is_heading:
            h_para = document.add_paragraph(style="Heading 2")
            h_para.paragraph_format.space_before = Pt(style.space_before_h2_pt)
            h_para.paragraph_format.space_after = Pt(style.space_after_h2_pt)
            h_para.paragraph_format.keep_with_next = True

            h_text = heading_title.upper() if style.all_caps_h2 else heading_title
            h_run = h_para.add_run(h_text)
            h_run.font.name = style.heading_font
            h_run.font.size = Pt(style.h2_size_pt)
            h_run.font.bold = True
            h_run.font.color.rgb = style.accent_rgb

            if style.add_heading_border:
                _add_bottom_border(h_para, style.accent_hex, style.border_size)

            idx += 1
            continue


        # 4. Bullet Points
        is_bullet, bullet_text = _is_bullet_point(stripped)
        if is_bullet:
            b_para = document.add_paragraph(style="List Bullet")
            b_para.paragraph_format.space_before = Pt(0)
            b_para.paragraph_format.space_after = Pt(2.5)
            b_para.paragraph_format.line_spacing = style.line_spacing

            _append_markdown_inline_runs(
                b_para,
                bullet_text,
                style.body_font,
                style.body_size_pt,
            )
            idx += 1
            continue

        # 5. Standard Body Paragraphs / Sub-Headings / Dates / Job Titles
        p_para = document.add_paragraph()
        p_para.paragraph_format.space_before = Pt(0)
        p_para.paragraph_format.space_after = Pt(style.space_after_body_pt)
        p_para.paragraph_format.line_spacing = style.line_spacing

        # Check if line looks like a job role or project title (e.g. bolded or contains company/dates)
        if stripped.startswith("**") or (" - " in stripped and any(c.isdigit() for c in stripped)):
            p_para.paragraph_format.space_before = Pt(2.5)

        _append_markdown_inline_runs(
            p_para,
            stripped,
            style.body_font,
            style.body_size_pt,
        )
        idx += 1

    # Save to temporary docx
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        temp_docx_path = tmp.name

    document.save(temp_docx_path)
    return temp_docx_path
