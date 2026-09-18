from pathlib import Path
import docx
import pytest

from services.docx_generator_service import (
    TEMPLATES,
    DEFAULT_TEMPLATE_ID,
    generate_resume_docx,
    get_available_templates,
)

SAMPLE_RESUME_MARKDOWN = """# Jane Doe
jane.doe@example.com | +1 (555) 019-2834 | San Francisco, CA | linkedin.com/in/janedoe

## Professional Summary
Accomplished Full Stack Software Engineer with 5+ years of experience designing scalable microservices, high-throughput REST APIs, and responsive web applications.

## Technical Skills
- **Languages**: Python, TypeScript, JavaScript, SQL
- **Frameworks**: FastAPI, React, Node.js, Next.js
- **Databases**: PostgreSQL, Redis, MongoDB
- **Cloud & DevOps**: AWS, Docker, Kubernetes, CI/CD, Git

## Professional Experience
**Senior Backend Engineer** | CloudSphere Inc. | 2022 - Present
- Architected and deployed event-driven microservices serving 20M+ daily requests with 99.99% uptime.
- Optimized PostgreSQL database indexes and query plans, cutting median endpoint latency by 45%.
- Led a team of 4 engineers transitioning monolithic architecture to containerized services on AWS ECS.

**Software Engineer** | AppWorks Studio | 2020 - 2022
- Developed customer-facing analytics dashboards with React and FastAPI used by 50,000+ monthly active users.
- Automated testing workflows with pytest and GitHub Actions, boosting test coverage from 60% to 92%.

## Education
**B.S. in Computer Science** | University of California, Berkeley | 2016 - 2020
- Graduated Magna Cum Laude (GPA: 3.85 / 4.0)
"""


def test_get_available_templates():
    templates = get_available_templates()
    assert len(templates) == 5
    template_ids = {t["template_id"] for t in templates}
    assert template_ids == {
        "modern_teal",
        "executive_navy",
        "tech_indigo",
        "elegant_burgundy",
        "compact_slate",
    }
    for t in templates:
        assert "name" in t
        assert "description" in t
        assert "persona" in t
        assert "accent_hex" in t
        assert "font_family" in t


@pytest.mark.parametrize("template_id", list(TEMPLATES.keys()))
def test_generate_all_templates(template_id):
    docx_path = generate_resume_docx(SAMPLE_RESUME_MARKDOWN, template_id=template_id)
    try:
        path = Path(docx_path)
        assert path.exists()
        assert path.stat().st_size > 1000

        # Verify it is a valid docx file that can be parsed
        doc = docx.Document(docx_path)
        assert len(doc.paragraphs) > 5

        # Verify title paragraph has candidate name
        title_para = doc.paragraphs[0]
        assert "Jane Doe" in title_para.text
    finally:
        Path(docx_path).unlink(missing_ok=True)


def test_generate_fallback_template():
    docx_path = generate_resume_docx(SAMPLE_RESUME_MARKDOWN, template_id="non_existent_template")
    try:
        path = Path(docx_path)
        assert path.exists()
        doc = docx.Document(docx_path)
        assert "Jane Doe" in doc.paragraphs[0].text
    finally:
        Path(docx_path).unlink(missing_ok=True)
