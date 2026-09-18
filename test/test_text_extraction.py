import os
import tempfile
from pathlib import Path
import pytest
from docx import Document
from fastapi import HTTPException, UploadFile

from services.text_extraction_service import (
    clean_text,
    extract_text_from_docx,
    extract_text_from_file,
    extract_text_from_txt,
)
from utils.file_utils import (
    RESUME_UPLOAD_EXTENSIONS,
    JOB_DESCRIPTION_UPLOAD_EXTENSIONS,
    _looks_like_heading,
    validate_upload_file,
    write_resume_docx,
)


def test_clean_text_normalizes_ligatures_and_quotes():
    raw = "ﬁnd the ﬂow with ‘curly quotes’ and “double quotes” — and – dashes."
    cleaned = clean_text(raw)
    assert "find the flow" in cleaned
    assert "'curly quotes'" in cleaned
    assert '"double quotes"' in cleaned
    assert "-- and - dashes" in cleaned or "- and - dashes" in cleaned


def test_clean_text_removes_bom_and_excess_spaces():
    raw = "\ufeff  Line 1   with   spaces.  \r\n\r\n\r\n\r\nLine 2.\u00a0\n"
    cleaned = clean_text(raw)
    assert "\ufeff" not in cleaned
    assert "Line 1 with spaces." in cleaned
    assert "Line 2." in cleaned
    assert "\n\n\n" not in cleaned


def test_extract_text_from_txt():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
        f.write("Candidate: John Doe\nSkill: Python, FastAPI\n")
        temp_path = f.name

    try:
        extracted = extract_text_from_txt(temp_path)
        assert "Candidate: John Doe" in extracted
        assert "Skill: Python, FastAPI" in extracted
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_extract_text_from_docx_with_tables():
    doc = Document()
    doc.add_paragraph("John Doe - Software Engineer")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Skills"
    table.cell(1, 0).text = "Backend"
    table.cell(1, 1).text = "Python, FastAPI, PostgreSQL"

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name

    try:
        doc.save(temp_path)
        extracted = extract_text_from_docx(temp_path)
        assert "John Doe - Software Engineer" in extracted
        assert "Category | Skills" in extracted or "Category" in extracted
        assert "Python, FastAPI, PostgreSQL" in extracted
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_looks_like_heading():
    assert _looks_like_heading("EXPERIENCE") is True
    assert _looks_like_heading("WORK EXPERIENCE") is True
    assert _looks_like_heading("EDUCATION:") is True
    assert _looks_like_heading("TECHNICAL SKILLS") is True
    assert _looks_like_heading("SUMMARY") is True
    assert _looks_like_heading("This is just a normal sentence explaining python experience.") is False
    assert _looks_like_heading("python django") is False


def test_write_resume_docx_creates_valid_docx():
    content = "JOHN DOE\n\nSUMMARY\nExperienced developer.\n\nEXPERIENCE\nBuilt scalable APIs."
    docx_path = write_resume_docx(content)
    try:
        assert Path(docx_path).exists()
        doc = Document(docx_path)
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "SUMMARY" in headings or "EXPERIENCE" in headings
    finally:
        Path(docx_path).unlink(missing_ok=True)


def test_validate_upload_file_rejects_doc():
    class DummyUpload:
        filename = "resume.doc"

    try:
        validate_upload_file(DummyUpload(), RESUME_UPLOAD_EXTENSIONS)
        assert False, "Expected HTTPException for .doc"
    except HTTPException as exc:
        assert exc.status_code == 400
        assert "Unsupported file type" in exc.detail


if __name__ == "__main__":
    test_clean_text_normalizes_ligatures_and_quotes()
    test_clean_text_removes_bom_and_excess_spaces()
    test_extract_text_from_txt()
    test_extract_text_from_docx_with_tables()
    test_looks_like_heading()
    test_write_resume_docx_creates_valid_docx()
    test_validate_upload_file_rejects_doc()
    print("ALL EXTRACTION TESTS PASSED!")
